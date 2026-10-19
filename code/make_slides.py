import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent.parent
bm = json.load(open(ROOT / "results/benchmark.json"))
an = json.load(open(ROOT / "results/analysis.json"))
m = bm["A. thaliana chloroplast"]["methods"]

INK, MUTED, ACC, BLUE, BG = RGBColor(0x1F, 0x29, 0x37), RGBColor(0x6B, 0x72, 0x80), RGBColor(0xD5, 0x5E, 0x00), RGBColor(0x00, 0x72, 0xB2), RGBColor(0xFA, 0xFA, 0xF8)
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]


def tb(sl, x, y, w, h, text, size=18, bold=False, color=INK, align=None):
    box = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = box.text_frame; tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, t in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = t; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color; p.space_after = Pt(8)
        if align: p.alignment = align
    return tf


def slide(title, n, notes=""):
    sl = prs.slides.add_slide(blank)
    sl.background.fill.solid(); sl.background.fill.fore_color.rgb = BG
    bar = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), prs.slide_height); bar.fill.solid(); bar.fill.fore_color.rgb = ACC; bar.line.fill.background()
    tb(sl, 0.6, 0.35, 12, 0.9, title, 30, True)
    tb(sl, 12.2, 6.95, 1, 0.4, f"{n}/10", 11, color=MUTED)
    sl.notes_slide.notes_text_frame.text = notes
    return sl


def bullets(sl, items, x=0.6, y=1.5, w=12, size=19):
    tf = tb(sl, x, y, w, 5, items, size)
    return tf


def table(sl, data, x, y, w, colw=None, size=12):
    rows, cols = len(data), len(data[0])
    t = sl.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(0.4 * rows)).table
    for i, r in enumerate(data):
        for j, v in enumerate(r):
            c = t.cell(i, j); c.text = str(v)
            for p in c.text_frame.paragraphs:
                p.font.size = Pt(size); p.font.bold = (i == 0)
    if colw:
        for j, cw in enumerate(colw): t.columns[j].width = Inches(cw)
    return t


# 1 Title
sl = prs.slides.add_slide(blank); sl.background.fill.solid(); sl.background.fill.fore_color.rgb = INK
tb(sl, 0.8, 1.8, 11.8, 2.2, "CodonRescore: Annotation-Free Codon-Model Rescoring of Ab Initio Gene Calls", 40, True, RGBColor(255, 255, 255))
tb(sl, 0.8, 4.0, 11.8, 1, "A benchmark on organellar and small prokaryotic genomes", 22, color=RGBColor(0xE6, 0x9F, 0x00))
tb(sl, 0.8, 5.4, 11.8, 1.2, ["Author 1 · Author 2 · Author 3", "Bioinformatics Research Work — Internal Assessment"], 16, color=RGBColor(0xD1, 0xD5, 0xDB))
sl.notes_slide.notes_text_frame.text = "Introduce topic: gene prediction is step one of genome annotation; we benchmark Pyrodigal and propose a lightweight rescoring step. Introduce team members and roles."

# 2 Problem & motivation
sl = slide("Problem and motivation", 2, "Errors in the CDS-finding stage propagate to all downstream annotation. Self-training needs enough sequence and may degrade on AT-rich or organellar genomes.")
bullets(sl, ["Gene prediction = first step of annotation; pipelines such as Bakta delegate CDS finding to Prodigal/Pyrodigal.",
             "Self-trained finders can degrade on short, AT-rich, RNA-gene-rich genomes (e.g. chloroplasts).",
             "Deep ORF classifiers (Balrog, ProtiGeno) need large pre-training; labs need a cheap, genome-specific check.",
             "Question: can a codon model trained only on the genome itself clean up a caller's output — and when does it fail?"])

# 3 Literature survey
sl = slide("Literature survey (recent work, 2021–2024)", 3, "Seven recent papers plus foundational references; see paper Table I.")
data = [["Work", "Approach", "Limitation for us"],
        ["Pyrodigal (2022)", "Prodigal port; self-training / meta models", "Needs ≥20 kb to train"],
        ["Balrog (2021)", "Universal protein model (TCN)", "Heavy pre-training"],
        ["Dimonaco et al. (2022)", "Benchmark of gene finders", "Organism-dependent accuracy"],
        ["Bakta (2021)", "Annotation pipeline on Prodigal CDS", "Inherits CDS errors"],
        ["ProtiGeno (2023)", "ESM-1b embeddings for short ORFs", "Lower precision; large model"],
        ["BRAKER3 (2024)", "Evidence integration (eukaryotes)", "Needs RNA-seq/protein data"],
        ["DNABERT-2 (2024)", "DNA foundation model", "Not tested for gene calling"]]
table(sl, data, 0.6, 1.5, 12.1, [3.0, 5.2, 3.9], 14)

# 4 Research gap
sl = slide("Research gap", 4, "Gaps are relative to the papers we reviewed, not a claim about all literature.")
bullets(sl, ["G1  Self-training on small organellar genomes is not examined in the works reviewed.",
             "G2  Cheap, genome-specific post-hoc rescoring (no pre-training, no annotation) is little explored.",
             "G3  Error anatomy of false positives is rarely reported — where should rescoring intervene?"], size=22)

# 5 Dataset
sl = slide("Datasets (real NCBI records)", 5, "All records are real RefSeq GenBank files; truth = single-exon, non-pseudo CDS. Spliced/pseudo/origin-spanning genes are neutral.")
data = [["Record", "Length", "GC", "CDS (all / eval.)", "Role"],
        ["A. thaliana chloroplast NC_000932", "154,478", "36.3%", "85 / 70", "Truth + training"],
        ["Y. pestis plasmid pPCP1 NC_005816", "9,609", "45.3%", "10 / 10", "Truth"],
        ["Phage phiX174 NC_001422", "5,386", "44.8%", "11 / 8", "Truth"],
        ["C. diphtheriae NCTC11397", "2,463,666", "53.5%", "no annotation", "Scalability only"]]
table(sl, data, 0.6, 1.7, 12.1, [4.6, 1.6, 1.1, 2.4, 2.4], 15)
tb(sl, 0.6, 4.4, 12, 1.5, "Scoring: stop-codon + strand match; precision / recall / F1; bootstrap 95% CI. Annotation is used for scoring only — never for training or tuning.", 17, color=MUTED)

# 6 Methodology
sl = slide("Methodology: CodonRescore", 6, "Three stages: seeds, model+threshold, rescoring. Parameters fixed a priori.")
bullets(sl, ["1  Enumerate ORFs (≥90 nt, both strands); seeds = longest, low-overlap ORFs (≥450 nt).",
             "2  3-periodic order-2 Markov coding model vs. genome background → per-codon log-likelihood ratio S.",
             "3  Threshold τ: 2-fold cross-fitted Youden cut-off, held-out seeds vs. shadow-frame ORFs.",
             "4  Re-score caller output: prune short calls with S<τ (<300 nt); rescue uncalled ORFs with S≥τ (≥120 nt).",
             "Applied to Pyrodigal single mode (self-trained) and meta mode (pre-trained)."], size=19)

# 7 Results
sl = slide("Results — chloroplast genome (70 evaluable genes)", 7, "Single-mode Pyrodigal over-calls (43 FP). CodonRescore: 43→16 FP, F1 0.678→0.789. Meta mode is strong (0.841) and CodonRescore does not improve it (0.816). CIs overlap.")
sl.shapes.add_picture(str(ROOT / "figs/fig_f1_bars.png"), Inches(0.5), Inches(1.4), height=Inches(4.2))
rows = [["Method", "Prec.", "Rec.", "F1"]]
for n, lab in (("Pyrodigal-single", "Pyrodigal single"), ("Pyrodigal-meta", "Pyrodigal meta"), ("CodonRescore(single)", "CodonRescore (single)"), ("CodonRescore(meta)", "CodonRescore (meta)")):
    s = m[n]; rows.append([lab, "%.3f" % s["precision"], "%.3f" % s["recall"], "%.3f" % s["f1"]])
table(sl, rows, 6.4, 1.6, 6.4, [3.1, 1.1, 1.1, 1.1], 15)
tb(sl, 6.4, 4.0, 6.4, 2.5, ["Single mode: 43 false positives → 16 after rescoring.",
                            "F1 95% CIs overlap (0.59–0.75 vs 0.72–0.85)."], 16)
tb(sl, 0.6, 5.8, 5.5, 0.8, "Also: pPCP1 and phiX174 — Pyrodigal-meta best; CodonRescore not applicable on phiX174 (too few seeds).", 13, color=MUTED)

# 8 Critical analysis
sl = slide("Critical analysis", 8, "Honest reading: partial repair of a weak caller, no gain over a strong one; small sample.")
bullets(sl, ["Of 43 single-mode false positives: 38 < 300 nt, 24 overlap tRNA/rRNA genes, none shadow a true protein gene.",
             "All methods find every gene ≥600 nt; differences are in the 25 genes <300 nt.",
             "Rescue step hurts: 0 of 10 rescued ORFs true in meta mode; order-2 model is too weak for rescue.",
             "Threshold is brittle: lowering τ by 0.1 drops F1 below 0.44; conservative is safer.",
             "Only 70 truth genes → differences mostly not significant; no hyper-parameter tuning was done."], size=19)

# 9 Proposed extension
sl = slide("Proposed research extension", 9, "Future work: more genomes, richer models, protein-level evidence, shared benchmark with Balrog/ProtiGeno.")
bullets(sl, ["Validate on many annotated organelle genomes + bacterial chromosomes for statistical power.",
             "Richer scorer: hexamer / higher-order Markov, start-site and RBS features.",
             "Protein-level evidence: protein / DNA language-model embeddings (ESM-1b, DNABERT-2) as rescoring features.",
             "Calibrated classifier instead of a single Youden cut-off; prune-only variant.",
             "Head-to-head with Balrog and ProtiGeno; extend to spliced genes and eukaryotic pipelines (BRAKER3 ideas)."], size=20)

# 10 Conclusion
sl = slide("Conclusion", 10, "Take-home message and team contributions.")
bullets(sl, ["Benchmarked Pyrodigal on real organellar / small genomes against RefSeq annotation.",
             "CodonRescore (annotation-free) cuts false positives of self-trained Pyrodigal 43 → 16 (F1 0.678 → 0.789).",
             "No improvement over pre-trained meta mode (0.816 vs 0.841) — and evidence is limited by 70 genes.",
             "Practical advice: prefer pre-trained/conservative settings on organelle genomes; distrust short, low-scoring calls near RNA genes.",
             "Manuscripts prepared in IEEE and Springer LNCS formats; code and results provided."], size=19)

prs.save(ROOT / "slides/Gene_Prediction_Presentation.pptx")
print("saved")
